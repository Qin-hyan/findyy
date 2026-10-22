# -*- coding: utf-8 -*-
"""Build the Findyy market research DOCX from report.md."""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK

BASE = r"D:\Findyy"
MD_PATH = os.path.join(BASE, "report.md")
OUT_PATH = os.path.join(BASE, "Findyy_市场调研报告_2026-10.docx")

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
GRAY = RGBColor(0x6B, 0x72, 0x80)
DARK = RGBColor(0x1F, 0x29, 0x37)
RED = RGBColor(0xC0, 0x39, 0x2B)

CN_BODY = "宋体"
CN_HEAD = "微软雅黑"


def set_cn_font(run, name=CN_BODY, size=10.5, bold=None, color=None, italic=None):
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(el, fill):
    tcPr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{k}")
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("第 ")
    set_cn_font(run, CN_BODY, 8.5, color=GRAY)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    r._r.append(fld1); r._r.append(instr); r._r.append(fld2); r._r.append(t); r._r.append(fld3)
    set_cn_font(r, CN_BODY, 8.5, color=GRAY)
    run2 = p.add_run(" 页")
    set_cn_font(run2, CN_BODY, 8.5, color=GRAY)


def add_header(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Findyy 市场调研  ·  2026-10  ·  内部研究")
    set_cn_font(run, CN_BODY, 8, color=GRAY)


def add_bottom_border(paragraph, color="1F3A5F", size=12):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_toc_field(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar = OxmlElement("w:fldChar"); fldChar.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve")
    instrText.text = 'TOC \\o "1-1" \\h \\z \\u'
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "目录将在打开文档时生成；如未显示，请右键“更新域”。"
    fldChar3 = OxmlElement("w:fldChar"); fldChar3.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar); run._r.append(instrText); run._r.append(fldChar2); run._r.append(t); run._r.append(fldChar3)
    set_cn_font(run, CN_BODY, 10, color=GRAY)


def add_cover(doc):
    # top navy band
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Cm(16.6)
    shade(cell._tc, "1F3A5F")
    set_cell_margins(cell, top=500, bottom=500, left=300, right=300)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Findyy")
    set_cn_font(r, CN_HEAD, 30, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    p2 = cell.add_paragraph()
    r2 = p2.add_run("桌面物品记忆相机")
    set_cn_font(r2, CN_HEAD, 20, bold=True, color=RGBColor(0xE9, 0xC4, 0x6A))

    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    r = p.add_run("市场调研与需求验证报告")
    set_cn_font(r, CN_HEAD, 26, bold=True, color=NAVY)

    p = doc.add_paragraph()
    r = p.add_run("竞品扫描 · 短视频需求证据 · 市场规模 · Go/No-Go 建议")
    set_cn_font(r, CN_HEAD, 12, color=GRAY)

    doc.add_paragraph()
    p = doc.add_paragraph()
    add_bottom_border(p, "2A9D8F", 18)

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    r = p.add_run("2026 年 10 月")
    set_cn_font(r, CN_HEAD, 14, bold=True, color=NAVY)

    p = doc.add_paragraph()
    r = p.add_run("数据截止 2026-10-22  ·  内部研究文档  ·  请勿外传")
    set_cn_font(r, CN_BODY, 9.5, color=GRAY)

    doc.add_page_break()


def add_disclaimer(doc):
    p = doc.add_paragraph()
    r = p.add_run("阅读说明")
    set_cn_font(r, CN_HEAD, 18, bold=True, color=NAVY)
    add_bottom_border(p, "1F3A5F", 8)

    items = [
        ("数据来源", "本报告仅使用公开网页、官方发布、公开电商/众筹页面与公开专利数据库；未进行一手问卷或深度访谈。所有关键数字均标注来源与采集日期；找不到可靠数据的项写“未找到/未核实”，不做估算冒充数据。"),
        ("证据分级", "A 级＝页面直接渲染或官方公报；B 级＝搜索摘要/媒体转述；C 级＝作者基于 A/B 级材料的归类与推断。正文中的评分、优先级、情景假设均属 C 级判断。"),
        ("竞品口径", "主竞品池只收“已量产在售/可实际购买/有真实用户评价”的产品；众筹未交付、研究原型、已停售产品单列“早期信号/失败案例”，不与在售产品混排。"),
        ("市场测算", "TAM/SAM/SOM 为方向性测算，不是预测。所有假设在正文与附录中显性列出；读者应重点关注假设的敏感性，而不是单一数字。"),
        ("汇率与日期", "美元按 1 USD≈7.15 CNY、港元按 1 HKD≈0.91 CNY、欧元按 1 EUR≈7.8 CNY 粗略换算；竞品状态与价格截至 2026-10-22，之后可能变化。"),
        ("非法律/投资意见", "隐私、合规、专利部分为产品风险提示，不构成法律意见；市场判断不构成投资建议。"),
    ]
    for title, text in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(f"{title}：")
        set_cn_font(r, CN_HEAD, 10.5, bold=True, color=NAVY)
        r2 = p.add_run(text)
        set_cn_font(r2, CN_BODY, 10, color=DARK)

    doc.add_page_break()


def add_toc_page(doc):
    p = doc.add_paragraph()
    r = p.add_run("目录")
    set_cn_font(r, CN_HEAD, 18, bold=True, color=NAVY)
    add_bottom_border(p, "1F3A5F", 8)
    add_toc_field(doc)
    doc.add_page_break()


def add_runs_with_format(paragraph, text, size=10.5, color=DARK, bold=False, font=CN_BODY):
    # split citations [S01] and bold **...**
    tokens = re.split(r"(\*\*.+?\*\*|\[S\d{2}\])", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = paragraph.add_run(tok[2:-2])
            set_cn_font(r, font, size, bold=True, color=color)
        elif re.fullmatch(r"\[S\d{2}\]", tok):
            r = paragraph.add_run(tok)
            set_cn_font(r, font, 8, color=GRAY)
        else:
            r = paragraph.add_run(tok)
            set_cn_font(r, font, size, bold=bold, color=color)


def add_paragraph(doc, text, style="body"):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if style == "body":
        pf.first_line_indent = Cm(0.74)
        pf.line_spacing = 1.35
        pf.space_after = Pt(3)
        add_runs_with_format(p, text, 10.5)
    elif style == "note":
        pf.first_line_indent = Cm(0)
        pf.line_spacing = 1.15
        pf.space_after = Pt(2)
        add_runs_with_format(p, text, 8.5, color=GRAY)
    elif style == "bullet":
        pf.left_indent = Cm(0.74)
        pf.first_line_indent = Cm(-0.37)
        pf.line_spacing = 1.3
        pf.space_after = Pt(1.5)
        r = p.add_run("• ")
        set_cn_font(r, CN_BODY, 10.5, color=TEAL, bold=True)
        add_runs_with_format(p, text, 10.5)
    elif style == "numbered":
        pf.left_indent = Cm(0.74)
        pf.first_line_indent = Cm(-0.37)
        pf.line_spacing = 1.3
        pf.space_after = Pt(1.5)
        add_runs_with_format(p, text, 10.5)
    return p


def add_quote(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.right_indent = Cm(0.6)
    pf.space_before = Pt(4)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.3
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F1F5F9"); pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18"); left.set(qn("w:space"), "8"); left.set(qn("w:color"), "2A9D8F")
    pBdr.append(left); pPr.append(pBdr)
    add_runs_with_format(p, text, 10, color=RGBColor(0x33, 0x41, 0x55))


def add_heading(doc, text, level, first_h1=False):
    p = doc.add_heading(level=level)
    if level == 1 and not first_h1:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(text)
        set_cn_font(r, CN_HEAD, 18, bold=True, color=NAVY)
        add_bottom_border(p, "1F3A5F", 8)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        set_cn_font(r, CN_HEAD, 13.5, bold=True, color=NAVY)
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_cn_font(r, CN_HEAD, 11.5, bold=True, color=RGBColor(0x2A, 0x5A, 0x7A))


def add_image(doc, path, alt):
    if not os.path.isabs(path):
        path = os.path.join(BASE, path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(path, width=Cm(16.0))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(alt)
    set_cn_font(r, CN_BODY, 8.5, color=GRAY)


def add_table(doc, headers, rows):
    ncols = len(headers)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Cm(16.0 / ncols)
        shade(cell._tc, "1F3A5F")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_runs_with_format(p, h, 8.3, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True, font=CN_HEAD)

    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.width = Cm(16.0 / ncols)
            set_cell_margins(cell)
            if ri % 2 == 1:
                shade(cell._tc, "F4F7FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            add_runs_with_format(p, text, 8.2, color=DARK)

    # repeat header row
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines):
    headers = [c.strip() for c in lines[0].strip().strip("|").split("|")]
    rows = []
    for line in lines[1:]:
        if re.match(r"^\|?\s*:?-{2,}", line.strip()):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return headers, rows


def build():
    doc = Document()
    # page setup A4
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)

    # Normal style
    st = doc.styles["Normal"]
    st.font.name = CN_BODY
    st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(4)

    # Heading styles
    for name, size, color in [("Heading 1", 18, NAVY), ("Heading 2", 13.5, NAVY), ("Heading 3", 11.5, RGBColor(0x2A, 0x5A, 0x7A))]:
        hs = doc.styles[name]
        hs.font.name = CN_HEAD
        hs._element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD)
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = color

    add_header(sec)
    add_page_number_footer(sec)

    # core properties
    cp = doc.core_properties
    cp.title = "Findyy 桌面物品记忆相机 · 市场调研与需求验证报告"
    cp.author = "Findyy Research"
    cp.subject = "市场调研、竞品分析、需求验证"
    cp.keywords = "物品记忆相机, AI camera, AirTag, 市场调研, TAM SAM SOM"

    add_cover(doc)
    add_disclaimer(doc)
    add_toc_page(doc)

    with open(MD_PATH, "r", encoding="utf-8") as f:
        lines = f.readlines()

    i = 0
    first_h1 = True
    in_code = False
    while i < len(lines):
        line = lines[i].rstrip("\n")
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            r = p.add_run(line)
            set_cn_font(r, "Consolas", 8.5, color=DARK)
            i += 1
            continue
        if s == "---":
            # Horizontal rule in the Markdown; H1 already handles page breaks.
            i += 1
            continue
        if s.startswith("# "):
            add_heading(doc, s[2:].strip(), 1, first_h1=first_h1)
            first_h1 = False
            i += 1
            continue
        if s.startswith("## "):
            add_heading(doc, s[3:].strip(), 2)
            i += 1
            continue
        if s.startswith("### "):
            add_heading(doc, s[4:].strip(), 3)
            i += 1
            continue
        if s.startswith("!["):
            m = re.match(r"!\[(.+?)\]\((.+?)\)", s)
            if m:
                add_image(doc, m.group(2), m.group(1))
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].rstrip("\n"))
                i += 1
            if len(block) >= 2:
                headers, rows = parse_table(block)
                add_table(doc, headers, rows)
            continue
        if s.startswith(">"):
            add_quote(doc, s[1:].strip())
            i += 1
            continue
        if s.startswith("- "):
            add_paragraph(doc, s[2:].strip(), "bullet")
            i += 1
            continue
        if re.match(r"^\d+\.\s", s):
            add_paragraph(doc, s, "numbered")
            i += 1
            continue
        if s.startswith("注：") or s.startswith("来源："):
            add_paragraph(doc, s, "note")
            i += 1
            continue
        add_paragraph(doc, s, "body")
        i += 1

    doc.save(OUT_PATH)
    print("saved", OUT_PATH)


if __name__ == "__main__":
    build()
